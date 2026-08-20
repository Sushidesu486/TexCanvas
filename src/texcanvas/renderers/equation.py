from __future__ import annotations

import re
from copy import deepcopy

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..mathml import latex_to_omml
from ..model import Slide
from .common import RenderContext, add_box, set_run_font

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


# Unicode replacements for common LaTeX math symbols and operators.
_LATEX_SYMBOLS: dict[str, str] = {
    "sum": "\u2211",        # ∑
    "prod": "\u220F",       # ∏
    "int": "\u222B",        # ∫
    "oint": "\u222E",       # ∮
    "partial": "\u2202",    # ∂
    "nabla": "\u2207",      # ∇
    "infty": "\u221E",      # ∞
    "cdot": "\u00B7",       # ·
    "times": "\u00D7",      # ×
    "div": "\u00F7",        # ÷
    "pm": "\u00B1",         # ±
    "mp": "\u2213",         # ∓
    "leq": "\u2264",        # ≤
    "geq": "\u2265",        # ≥
    "neq": "\u2260",        # ≠
    "approx": "\u2248",     # ≈
    "equiv": "\u2261",      # ≡
    "rightarrow": "\u2192", # →
    "leftarrow": "\u2190",  # ←
    "Rightarrow": "\u21D2", # ⇒
    "Leftarrow": "\u21D0",  # ⇐
    "leftrightarrow": "\u2194",  # ↔
    "in": "\u2208",         # ∈
    "notin": "\u2209",      # ∉
    "subset": "\u2282",     # ⊂
    "supset": "\u2283",     # ⊃
    "cup": "\u222A",        # ∪
    "cap": "\u2229",        # ∩
    "forall": "\u2200",     # ∀
    "exists": "\u2203",    # ∃
    "neg": "\u00AC",        # ¬
    "land": "\u2227",       # ∧
    "lor": "\u2228",        # ∨
    "to": "\u2192",         # →
    "sqrt": "\u221A",       # √
    "alpha": "\u03B1", "beta": "\u03B2", "gamma": "\u03B3", "delta": "\u03B4",
    "epsilon": "\u03B5", "varepsilon": "\u03B5", "zeta": "\u03B6", "eta": "\u03B7",
    "theta": "\u03B8", "vartheta": "\u03D1", "iota": "\u03B9", "kappa": "\u03BA",
    "lambda": "\u03BB", "mu": "\u03BC", "nu": "\u03BD", "xi": "\u03BE",
    "pi": "\u03C0", "varpi": "\u03D6", "rho": "\u03C1", "varrho": "\u03F1",
    "sigma": "\u03C3", "varsigma": "\u03C2", "tau": "\u03C4", "upsilon": "\u03C5",
    "phi": "\u03C6", "varphi": "\u03D5", "chi": "\u03C7", "psi": "\u03C8",
    "omega": "\u03C9", "Gamma": "\u0393", "Delta": "\u0394", "Theta": "\u0398",
    "Lambda": "\u039B", "Xi": "\u039E", "Pi": "\u03A0", "Sigma": "\u03A3",
    "Phi": "\u03A6", "Psi": "\u03A8", "Omega": "\u03A9",
}

# Commands that render as upright roman text (function names) rather than symbols.
_LATEX_FUNCTIONS = {
    "log", "ln", "exp", "sin", "cos", "tan", "csc", "sec", "cot",
    "arcsin", "arccos", "arctan", "sinh", "cosh", "tanh", "lim", "max", "min",
    "sup", "inf", "arg", "deg", "det", "dim", "ker", "hom", "Pr",
}

_TOKEN_PATTERN = re.compile(
    r"\\frac\{([^{}]*)\}\{([^{}]*)\}"        # \frac{a}{b}
    r"|\\([A-Za-z]+)"                         # \command
    r"|\^(\{[^{}]*\}|.)"                       # ^{...} or ^x
    r"|_(\{[^{}]*\}|.)"                       # _{...} or _x
)


def _add_run(paragraph, text: str, *, latin: str, ea: str, size: float, text_color: str, baseline: str | None = None) -> None:
    run = paragraph.add_run()
    run.text = text
    set_run_font(run, latin=latin, ea=ea)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(text_color)
    if baseline is not None:
        run.font._rPr.set("baseline", baseline)


def _add_script(paragraph, token: str, latin: str, ea: str, size: float, text_color: str, *, superscript: bool) -> None:
    text = token[1:-1] if token.startswith("{") and token.endswith("}") else token
    baseline = "30000" if superscript else "-25000"
    # Sub/superscripts render slightly smaller.
    _add_run(paragraph, text, latin=latin, ea=ea, size=size * 0.75, text_color=text_color, baseline=baseline)


def _render_inline(value: str, paragraph, *, latin: str, ea: str, size: float, text_color: str) -> None:
    """Add styled runs, applying minimal LaTeX-style math markup.

    Supports ``\\frac{a}{b}``, ``\\sum``/``\\alpha``/... symbol commands,
    ``\\log``/``\\sin``/... function names, ``^{...}`` superscripts and
    ``_{...}`` subscripts. Unrecognized text renders as plain runs.
    """
    cursor = 0
    has_match = False
    for match in _TOKEN_PATTERN.finditer(value):
        has_match = True
        if match.start() > cursor:
            _add_run(paragraph, value[cursor:match.start()], latin=latin, ea=ea, size=size, text_color=text_color)
        if match.group(1) is not None and match.group(2) is not None:
            # \frac{a}{b} -> a⁄b using a fraction slash, kept inline for portability.
            _add_run(paragraph, match.group(1), latin=latin, ea=ea, size=size * 0.8, text_color=text_color)
            _add_run(paragraph, "\u2044", latin=latin, ea=ea, size=size, text_color=text_color)
            _add_run(paragraph, match.group(2), latin=latin, ea=ea, size=size * 0.8, text_color=text_color)
        elif match.group(3) is not None:
            command = match.group(3)
            if command in _LATEX_SYMBOLS:
                _add_run(paragraph, _LATEX_SYMBOLS[command], latin=latin, ea=ea, size=size, text_color=text_color)
            elif command in _LATEX_FUNCTIONS:
                _add_run(paragraph, command, latin=latin, ea=ea, size=size, text_color=text_color)
            else:
                # Unknown command: drop the backslash and render the name literally.
                _add_run(paragraph, command, latin=latin, ea=ea, size=size, text_color=text_color)
        elif match.group(4) is not None:
            _add_script(paragraph, match.group(4), latin, ea, size, text_color, superscript=True)
        elif match.group(5) is not None:
            _add_script(paragraph, match.group(5), latin, ea, size, text_color, superscript=False)
        cursor = match.end()
    if not has_match:
        _add_run(paragraph, value, latin=latin, ea=ea, size=size, text_color=text_color)
    elif cursor < len(value):
        _add_run(paragraph, value[cursor:], latin=latin, ea=ea, size=size, text_color=text_color)


def render_equation(ctx: RenderContext, slide: Slide) -> None:
    assert slide.equation
    equation = slide.equation.strip("\n")
    display_lines = [line for line in equation.split("\n") if line.strip()]

    panel_top = 1.6
    available_height = ctx.slide_height - 1.48 - 0.6 - panel_top
    panel_height = min(available_height, max(1.2, 0.6 * len(display_lines) + 0.6))
    panel_width = ctx.slide_width - 2.4
    panel_left = (ctx.slide_width - panel_width) / 2

    add_box(
        ctx,
        "DSH_EQUATION_PANEL",
        panel_left,
        panel_top,
        panel_width,
        panel_height,
        fill=ctx.theme.equation_fill,
        line=ctx.theme.equation_border,
        radius=False,
    )

    shape = ctx.slide.shapes.add_textbox(
        Inches(panel_left + 0.2),
        Inches(panel_top + 0.15),
        Inches(panel_width - 0.4),
        Inches(max(0.4, panel_height - 0.3)),
    )
    shape.name = "DSH_EQUATION"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = frame.margin_right = Inches(0.1)
    frame.margin_top = frame.margin_bottom = Inches(0.05)

    size = 22 if len(display_lines) <= 2 else 18 if len(display_lines) <= 4 else 15

    for index, line in enumerate(display_lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.line_spacing = 1.2
        paragraph.space_after = Pt(8)
        paragraph.text = ""
        if not _render_omml_line(paragraph, line, latin=ctx.theme.latin_font, ea=ctx.theme.chinese_font):
            _render_inline(line, paragraph, latin=ctx.theme.latin_font, ea=ctx.theme.chinese_font, size=size, text_color=ctx.theme.primary)


def _render_omml_line(paragraph, latex: str, *, latin: str, ea: str) -> bool:
    """Try to render the line as a native OMML equation.

    Returns True when a pandoc-generated ``<m:oMath>`` was appended; False when
    pandoc is unavailable or the line is not a LaTeX equation, in which case
    the caller should fall back to the legacy Unicode renderer.

    The ``oMath`` is wrapped in an ``oMathPara`` (with centered justification)
    because PowerPoint/WPS render a block-level equation only when it sits
    inside ``m:oMathPara`` — a bare ``m:oMath`` directly under ``a:p`` is
    treated as inline math and some renderers show it as empty.

    ``latin``/``ea`` are applied to the math runs' ``<m:rPr>`` so the equation
    honors the deck's font rules (Latin Helvetica, East-Asian 苹方-简).
    """
    omath = latex_to_omml(latex)
    if omath is None:
        return False
    _apply_fonts_to_math(omath, latin=latin, ea=ea)
    oMathPara = etree.SubElement(paragraph._p, "{%s}oMathPara" % MATH_NS)
    oMathParaPr = etree.SubElement(oMathPara, "{%s}oMathParaPr" % MATH_NS)
    jc = etree.SubElement(oMathParaPr, "{%s}jc" % MATH_NS)
    jc.set("{%s}val" % MATH_NS, "center")
    oMathPara.append(deepcopy(omath))
    return True


def _apply_fonts_to_math(omath, *, latin: str, ea: str) -> None:
    """Stamp ``a:latin``/``a:ea``/``a:cs`` into every ``<m:rPr>`` (or create one).

    Without this, PowerPoint/WPS fall back to the theme math font (Cambria Math)
    for Latin and the theme East-Asian font for CJK glyphs inside the equation.
    """
    for run in omath.iter("{%s}r" % MATH_NS):
        rPr = run.find("{%s}rPr" % MATH_NS)
        if rPr is None:
            rPr = etree.SubElement(run, "{%s}rPr" % MATH_NS)
            # rPr must precede the <m:t> child per the schema ordering.
            run.insert(0, rPr)
        for tag, typeface in (("a:latin", latin), ("a:ea", ea), ("a:cs", ea)):
            existing = rPr.find("{%s}%s" % (A_NS, tag.split(":")[1]))
            if existing is None:
                existing = etree.SubElement(rPr, "{%s}%s" % (A_NS, tag.split(":")[1]))
            existing.set("typeface", typeface)
